#!/usr/bin/env python3
"""Regenerate store-master.json from the store master workbook.

    python update-store-master.py "C:\\Users\\dgautam\\OneDrive - Kuwait Food Company\\Store Master Sheet - July'26 (1).xlsx"

Reads the "For Loyalty" sheet (Store ID 135xxx, Rest. ID 35xxx) and takes the cleaner
store names and menu type from "For HD" where the Rest. ID matches. Only store identity
columns are copied: no phone numbers, emails or staff names.
"""

import datetime
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).with_name("store-master.json")


def text(value):
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip() if value is not None else ""


def sheet_rows(workbook, name):
    rows = workbook[name].iter_rows(values_only=True)
    header = [text(h) for h in next(rows)]
    for row in rows:
        if any(v not in (None, "") for v in row):
            yield {header[i]: row[i] for i in range(min(len(header), len(row))) if header[i]}


def main(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    hd = {text(r.get("REST. ID")): r for r in sheet_rows(workbook, "For HD")}
    stores = []
    for r in sheet_rows(workbook, "For Loyalty"):
        store_id, rest_id = text(r.get("STORE ID")), text(r.get("REST. ID"))
        if not store_id:
            continue
        h = hd.get(rest_id, {})
        stores.append({
            "storeId": store_id,
            "restId": rest_id,
            "name": text(h.get("STORE NAME")) or text(r.get("STORE NAME")),
            "display": text(r.get("DISPLAY NAME")),
            "emirate": text(r.get("EMIRATE")),
            "opened": text(r.get("DOO")),
            "menu": text(h.get("MENU TYPE")) or text(r.get("MENU TYPE")),
            "pricing": text(r.get("PRICING")),
            "inHelpdeskSheet": bool(h),
        })
    stores.sort(key=lambda s: s["storeId"])
    OUT.write_text(json.dumps({
        "source": Path(path).name,
        "sheet": "For Loyalty",
        "extracted": datetime.date.today().isoformat(),
        "stores": stores,
    }, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"{len(stores)} stores written to {OUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
